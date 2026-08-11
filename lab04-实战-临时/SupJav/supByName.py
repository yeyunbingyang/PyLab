
import sys
import os
import re
import time
import io
from datetime import datetime
from dataclasses import dataclass
from typing import List, Set, Dict, Optional

import requests
from PIL import Image as PILImage

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QTextEdit, QProgressBar,
    QFileDialog, QGroupBox, QSpinBox, QCheckBox, QMessageBox,
    QTableWidget, QTableWidgetItem, QHeaderView, QSplitter,
    QStatusBar, QMenuBar, QMenu
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QSettings, QUrl
from PyQt6.QtGui import QColor, QBrush, QFont, QAction, QDesktopServices
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font as OpenpyxlFont, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from DrissionPage import Chromium


# ==================== 配色常量（暗色主题）====================
DARK_BG = "#1e1e2e"
DARK_CARD = "#2a2a3c"
DARK_INPUT = "#313244"
TEXT_PRIMARY = "#cdd6f4"
TEXT_SECONDARY = "#a6adc8"
ACCENT_GREEN = "#a6e3a1"
ACCENT_RED = "#f38ba8"
ACCENT_BLUE = "#89b4fa"
ACCENT_YELLOW = "#f9e2af"
ACCENT_PURPLE = "#cba6f7"
BORDER_COLOR = "#45475a"


# ==================== 数据模型 ====================
@dataclass
class VideoInfo:
    code: str
    title: str
    url: str
    img_url: str
    update_date: str
    views: str
    is_local: bool = False


# ==================== 工具函数 ====================
def parse_meta_text(meta_text: str, meta_html: str = "") -> tuple:
    update_date, views = "", ""
    txt = meta_text.strip()
    if not txt:
        return update_date, views
    lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
    if len(lines) >= 2:
        update_date = lines[0]
        views = lines[1]
        return update_date, views
    if len(lines) == 1:
        line = lines[0]
        m = re.search(r'(\d{4}/\d{2}/\d{2})(\d+[\s\w]*)', line)
        if m:
            update_date = m.group(1)
            views = m.group(2).strip()
            if not re.search(r'Views', views, re.I):
                views_match = re.search(r'(\d+\s*Views)', line, re.I)
                if views_match:
                    views = views_match.group(1)
            return update_date, views
        if meta_html:
            date_m = re.search(r'(\d{4}/\d{2}/\d{2})', meta_html)
            if date_m:
                update_date = date_m.group(1)
            views_m = re.search(r'(\d+)\s*Views', meta_html, re.I)
            if views_m:
                views = views_m.group(0)
            if update_date or views:
                return update_date, views
        update_date = line
    return update_date, views


def download_image(url: str, max_size: tuple = (120, 160)) -> Optional[bytes]:
    if not url:
        return None
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Referer": "https://supjav.com/"
        }
        resp = requests.get(url, headers=headers, timeout=10)
        resp.raise_for_status()
        img = PILImage.open(io.BytesIO(resp.content))
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.thumbnail(max_size, PILImage.Resampling.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return None


# ==================== 工作线程 ====================
class ScrapyWorker(QThread):
    log_signal = pyqtSignal(str)
    progress_signal = pyqtSignal(int, int)
    data_signal = pyqtSignal(list)
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, name: str, save_dir: str, browser_port: int = 9333,
                 embed_images: bool = True):
        super().__init__()
        self.name = name
        self.save_dir = save_dir
        self.browser_port = browser_port
        self.embed_images = embed_images
        self._is_running = True
        self.browser = None
        self.tab = None

    def log(self, msg: str):
        self.log_signal.emit(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}")

    def stop(self):
        self._is_running = False
        self.log("正在停止...")

    def run(self):
        try:
            self.log(f"开始搜索: {self.name}")
            self.browser = Chromium(self.browser_port)
            self.tab = self.browser.new_tab()
            local_codes = self._get_local_codes()
            self.log(f"本地已有番号: {len(local_codes)} 个")
            raw_data = self._fetch_all_data()
            if not self._is_running:
                return
            self.log(f"原始数据: {len(raw_data)} 条")
            deduped_data = self._deduplicate_by_code(raw_data)
            self.log(f"去重后: {len(deduped_data)} 条")
            for item in deduped_data:
                item.is_local = item.code in local_codes
            local_count = sum(1 for item in deduped_data if item.is_local)
            self.log(f"本地已存在: {local_count} 条")
            if deduped_data and self._is_running:
                filepath = self._save_to_excel(deduped_data)
                self.data_signal.emit(deduped_data)
                self.finished_signal.emit(True, f"完成！已保存到: {filepath}")
            else:
                self.finished_signal.emit(False, "没有抓取到数据")
        except Exception as e:
            self.log(f"错误: {str(e)}")
            self.finished_signal.emit(False, f"发生错误: {str(e)}")
        finally:
            if self.tab:
                try:
                    self.tab.close()
                except:
                    pass

    def _get_local_codes(self) -> Set[str]:
        codes = set()
        pattern = re.compile(r'[A-Z]{2,8}-\d{2,5}', re.I)
        if not self.save_dir or not os.path.exists(self.save_dir):
            return codes
        for root, _, files in os.walk(self.save_dir):
            for file in files:
                m = pattern.search(file)
                if m:
                    codes.add(m.group().upper())
        return codes

    def _fetch_all_data(self) -> List[VideoInfo]:
        self.tab.get(f'https://supjav.com/ja/?s={self.name}')
        time.sleep(2)
        lis = self.tab.eles('xpath://ul/li')
        page_count = 1
        for li in lis:
            text = li.text.strip()
            if text.isdigit():
                page_count = max(page_count, int(text))
        self.log(f"总页数: {page_count}")
        all_posts = []
        for page in range(1, page_count + 1):
            if not self._is_running:
                break
            self.log(f"正在抓取第 {page}/{page_count} 页...")
            self.progress_signal.emit(page, page_count)
            self.tab.get(f'https://supjav.com/ja/page/{page}?s={self.name}')
            time.sleep(1.5)
            posts = self.tab.eles('css:div.post')
            for post in posts:
                if not self._is_running:
                    break
                a = post.ele('css:h3 a')
                if not a:
                    continue
                title = a.attr('title') or a.text
                if '[モザイク破壊]' not in title:
                    continue
                pattern = re.compile(r'[A-Z]{2,8}-\d{2,5}', re.I)
                m = pattern.search(title)
                code = m.group().upper() if m else ''
                url = a.attr('href')
                img = post.ele('css:img.thumb')
                img_url = ''
                if img:
                    img_url = img.attr('data-original') or img.attr('src') or ''
                meta = post.ele('css:div.meta')
                update_date, views = '', ''
                if meta:
                    meta_text = meta.text or ''
                    meta_html = getattr(meta, 'html', '') or ''
                    update_date, views = parse_meta_text(meta_text, meta_html)
                all_posts.append(VideoInfo(
                    code=code, title=title, url=url, img_url=img_url,
                    update_date=update_date, views=views
                ))
        return all_posts

    def _deduplicate_by_code(self, posts: List[VideoInfo]) -> List[VideoInfo]:
        code_map: Dict[str, VideoInfo] = {}
        for post in posts:
            if not post.code:
                continue
            try:
                current_date = datetime.strptime(post.update_date, "%Y/%m/%d")
            except (ValueError, TypeError):
                current_date = datetime.min
            if post.code not in code_map:
                code_map[post.code] = post
            else:
                try:
                    existing_date = datetime.strptime(code_map[post.code].update_date, "%Y/%m/%d")
                except (ValueError, TypeError):
                    existing_date = datetime.min
                if current_date > existing_date:
                    code_map[post.code] = post
        return list(code_map.values())

    def _save_to_excel(self, data: List[VideoInfo]) -> str:
        today = datetime.now().strftime("%Y%m%d")
        filename = f"{self.name}-{today}.xlsx"
        filepath = os.path.join(self.save_dir, filename) if self.save_dir else filename
        wb = Workbook()
        ws = wb.active
        ws.title = "搜索结果"
        header_fill = PatternFill(start_color="585b70", end_color="585b70", fill_type="solid")
        header_font = OpenpyxlFont(bold=True, color="cdd6f4", size=11)
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color="45475a"),
            right=Side(style='thin', color="45475a"),
            top=Side(style='thin', color="45475a"),
            bottom=Side(style='thin', color="45475a")
        )
        yellow_fill = PatternFill(start_color="f9e2af", end_color="f9e2af", fill_type="solid")
        yellow_font = OpenpyxlFont(color="1e1e2e")
        headers = ["番号", "标题", "URL", "图片", "更新日期", "查看量", "本地已存在"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border
        if self.embed_images:
            ws.row_dimensions[1].height = 25
        for item in data:
            row_idx = ws.max_row + 1
            ws.cell(row=row_idx, column=1, value=item.code)
            ws.cell(row=row_idx, column=2, value=item.title)
            ws.cell(row=row_idx, column=3, value=item.url)
            ws.cell(row=row_idx, column=5, value=item.update_date)
            ws.cell(row=row_idx, column=6, value=item.views)
            ws.cell(row=row_idx, column=7, value="是" if item.is_local else "否")
            if item.is_local:
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col)
                    cell.fill = yellow_fill
                    cell.font = yellow_font
                    cell.border = thin_border
            else:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col).border = thin_border
            if self.embed_images and item.img_url:
                self.log(f"正在下载图片: {item.code}")
                img_bytes = download_image(item.img_url, max_size=(120, 160))
                if img_bytes:
                    try:
                        img = OpenpyxlImage(io.BytesIO(img_bytes))
                        img.width = 100
                        img.height = 140
                        ws.add_image(img, f"D{row_idx}")
                        ws.row_dimensions[row_idx].height = 110
                    except Exception as e:
                        self.log(f"图片嵌入失败 {item.code}: {e}")
                        ws.cell(row=row_idx, column=4, value=item.img_url)
                else:
                    ws.cell(row=row_idx, column=4, value=item.img_url)
            else:
                ws.cell(row=row_idx, column=4, value=item.img_url)
        col_widths = {"A": 12, "B": 50, "C": 40, "D": 18, "E": 14, "F": 14, "G": 12}
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width
        ws.freeze_panes = "A2"
        wb.save(filepath)
        self.log(f"已保存: {filepath}")
        return filepath


# ==================== 主窗口 ====================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Supjav 爬虫工具")
        self.setMinimumSize(1300, 850)
        self.settings = QSettings("SupjavScraper", "Settings")
        self.worker: Optional[ScrapyWorker] = None
        self.current_data: List[VideoInfo] = []
        self._init_ui()
        self._load_settings()

    def _init_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(20, 20, 20, 20)

        # === 配置区域 ===
        config_group = QGroupBox("⚙️ 配置")
        config_group.setObjectName("configGroup")
        config_layout = QVBoxLayout(config_group)
        config_layout.setSpacing(10)

        name_layout = QHBoxLayout()
        name_label = QLabel("搜索名称:")
        name_label.setObjectName("label")
        name_layout.addWidget(name_label)
        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("例如: 美乃すずめ")
        self.name_input.setText("美乃すずめ")
        self.name_input.setObjectName("lineEdit")
        name_layout.addWidget(self.name_input)
        config_layout.addLayout(name_layout)

        dir_layout = QHBoxLayout()
        dir_label = QLabel("保存目录:")
        dir_label.setObjectName("label")
        dir_layout.addWidget(dir_label)
        self.dir_input = QLineEdit()
        self.dir_input.setPlaceholderText("选择Excel保存和本地视频所在目录...")
        self.dir_input.setReadOnly(True)
        self.dir_input.setObjectName("lineEdit")
        dir_layout.addWidget(self.dir_input)
        self.browse_btn = QPushButton("📁 浏览...")
        self.browse_btn.setObjectName("accentBtn")
        self.browse_btn.clicked.connect(self._browse_directory)
        dir_layout.addWidget(self.browse_btn)
        config_layout.addLayout(dir_layout)

        port_layout = QHBoxLayout()
        port_label = QLabel("浏览器端口:")
        port_label.setObjectName("label")
        port_layout.addWidget(port_label)
        self.port_spin = QSpinBox()
        self.port_spin.setRange(1024, 65535)
        self.port_spin.setValue(9333)
        self.port_spin.setObjectName("spinBox")
        port_layout.addWidget(self.port_spin)

        self.embed_img_check = QCheckBox("嵌入图片到Excel")
        self.embed_img_check.setChecked(True)
        self.embed_img_check.setObjectName("checkBox")
        port_layout.addWidget(self.embed_img_check)

        port_layout.addStretch()
        config_layout.addLayout(port_layout)

        main_layout.addWidget(config_group)

        # === 控制按钮 ===
        btn_layout = QHBoxLayout()
        btn_layout.setSpacing(12)

        self.start_btn = QPushButton("▶️ 开始抓取")
        self.start_btn.setObjectName("greenBtn")
        self.start_btn.setMinimumHeight(42)
        self.start_btn.clicked.connect(self._start_scraping)
        btn_layout.addWidget(self.start_btn)

        self.stop_btn = QPushButton("⏹️ 停止")
        self.stop_btn.setObjectName("redBtn")
        self.stop_btn.setMinimumHeight(42)
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self._stop_scraping)
        btn_layout.addWidget(self.stop_btn)

        self.export_btn = QPushButton("📊 导出Excel")
        self.export_btn.setObjectName("blueBtn")
        self.export_btn.setMinimumHeight(42)
        self.export_btn.setEnabled(False)
        self.export_btn.clicked.connect(self._export_data)
        btn_layout.addWidget(self.export_btn)

        btn_layout.addStretch()
        main_layout.addLayout(btn_layout)

        # === 进度条 ===
        self.progress_bar = QProgressBar()
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setFormat("第 %v/%m 页")
        self.progress_bar.setVisible(False)
        self.progress_bar.setObjectName("progressBar")
        main_layout.addWidget(self.progress_bar)

        # === 分割器 ===
        splitter = QSplitter(Qt.Orientation.Vertical)

        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "番号", "标题", "URL", "图片", "更新日期", "查看量", "本地已存在"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self._show_table_context_menu)
        self.table.setObjectName("dataTable")
        splitter.addWidget(self.table)

        log_group = QGroupBox("📝 日志")
        log_group.setObjectName("logGroup")
        log_layout = QVBoxLayout(log_group)
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setFont(QFont("JetBrains Mono", 10))
        self.log_text.setObjectName("logText")
        log_layout.addWidget(self.log_text)
        splitter.addWidget(log_group)

        splitter.setSizes([500, 300])
        main_layout.addWidget(splitter, 1)

        self.status_bar = QStatusBar()
        self.status_bar.setObjectName("statusBar")
        self.setStatusBar(self.status_bar)
        self.status_bar.showMessage("就绪")

        menubar = self.menuBar()
        menubar.setObjectName("menuBar")
        file_menu = menubar.addMenu("文件")
        file_menu.setObjectName("menu")
        exit_action = QAction("退出", self)
        exit_action.setShortcut("Ctrl+Q")
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

        help_menu = menubar.addMenu("帮助")
        help_menu.setObjectName("menu")
        about_action = QAction("关于", self)
        about_action.triggered.connect(self._show_about)
        help_menu.addAction(about_action)

    def _load_settings(self):
        saved_dir = self.settings.value("save_dir", "")
        if saved_dir and os.path.exists(saved_dir):
            self.dir_input.setText(saved_dir)
        saved_port = self.settings.value("browser_port", 9333)
        self.port_spin.setValue(int(saved_port))
        saved_name = self.settings.value("last_name", "美乃すずめ")
        self.name_input.setText(saved_name)
        embed = self.settings.value("embed_images", "true")
        self.embed_img_check.setChecked(embed.lower() == "true")

    def _save_settings(self):
        self.settings.setValue("save_dir", self.dir_input.text())
        self.settings.setValue("browser_port", self.port_spin.value())
        self.settings.setValue("last_name", self.name_input.text())
        self.settings.setValue("embed_images", str(self.embed_img_check.isChecked()).lower())

    def _browse_directory(self):
        current_dir = self.dir_input.text() or os.path.expanduser("~")
        directory = QFileDialog.getExistingDirectory(
            self, "选择保存目录", current_dir,
            QFileDialog.Option.ShowDirsOnly
        )
        if directory:
            self.dir_input.setText(directory)
            self._save_settings()

    def _start_scraping(self):
        name = self.name_input.text().strip()
        save_dir = self.dir_input.text().strip()
        port = self.port_spin.value()
        if not name:
            QMessageBox.warning(self, "警告", "请输入搜索名称！")
            return
        if not save_dir:
            reply = QMessageBox.question(
                self, "确认", "未选择保存目录，将使用当前目录。是否继续？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.No:
                return
            save_dir = os.getcwd()
        self._save_settings()
        self.table.setRowCount(0)
        self.log_text.clear()
        self.current_data.clear()
        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.export_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.status_bar.showMessage("正在抓取...")
        self.worker = ScrapyWorker(name, save_dir, port, self.embed_img_check.isChecked())
        self.worker.log_signal.connect(self._append_log)
        self.worker.progress_signal.connect(self._update_progress)
        self.worker.data_signal.connect(self._update_table)
        self.worker.finished_signal.connect(self._on_finished)
        self.worker.start()

    def _stop_scraping(self):
        if self.worker:
            self.worker.stop()
            self.worker.wait(5000)
        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        self.progress_bar.setVisible(False)
        self.status_bar.showMessage("已停止")

    def _update_progress(self, current: int, total: int):
        self.progress_bar.setMaximum(total)
        self.progress_bar.setValue(current)

    def _append_log(self, msg: str):
        self.log_text.append(msg)
        scrollbar = self.log_text.verticalScrollBar()
        scrollbar.setValue(scrollbar.maximum())

    def _update_table(self, data: List[VideoInfo]):
        self.current_data = data
        self.table.setRowCount(len(data))
        for row, item in enumerate(data):
            values = [
                item.code, item.title, item.url, item.img_url,
                item.update_date, item.views, "是" if item.is_local else "否"
            ]
            for col, value in enumerate(values):
                cell = QTableWidgetItem(str(value))
                if item.is_local:
                    cell.setBackground(QBrush(QColor(249, 226, 175)))  # 暗色主题下的黄色
                    cell.setForeground(QBrush(QColor(30, 30, 46)))     # 深色文字
                self.table.setItem(row, col, cell)
        self.table.resizeColumnsToContents()

    def _on_finished(self, success: bool, msg: str):
        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        self.export_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        if success:
            self.status_bar.showMessage(f"完成 - {msg}")
            QMessageBox.information(self, "完成", msg)
        else:
            self.status_bar.showMessage(f"失败 - {msg}")
            QMessageBox.critical(self, "错误", msg)

    def _export_data(self):
        if not self.current_data:
            QMessageBox.warning(self, "警告", "没有数据可导出！")
            return
        filepath, _ = QFileDialog.getSaveFileName(
            self, "保存Excel",
            f"{self.name_input.text()}-{datetime.now().strftime('%Y%m%d')}.xlsx",
            "Excel Files (*.xlsx)"
        )
        if filepath:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "搜索结果"
                headers = ["番号", "标题", "URL", "图片", "更新日期", "查看量", "本地已存在"]
                ws.append(headers)
                yellow_fill = PatternFill(start_color="f9e2af", end_color="f9e2af", fill_type="solid")
                yellow_font = OpenpyxlFont(color="1e1e2e")
                thin_border = Border(
                    left=Side(style='thin', color="45475a"),
                    right=Side(style='thin', color="45475a"),
                    top=Side(style='thin', color="45475a"),
                    bottom=Side(style='thin', color="45475a")
                )
                for item in self.current_data:
                    ws.append([
                        item.code, item.title, item.url, item.img_url,
                        item.update_date, item.views, "是" if item.is_local else "否"
                    ])
                    if item.is_local:
                        for col in range(1, 8):
                            cell = ws.cell(row=ws.max_row, column=col)
                            cell.fill = yellow_fill
                            cell.font = yellow_font
                            cell.border = thin_border
                wb.save(filepath)
                QMessageBox.information(self, "成功", f"已导出到: {filepath}")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"导出失败: {str(e)}")

    def _show_table_context_menu(self, position):
        menu = QMenu()
        menu.setObjectName("contextMenu")
        copy_action = menu.addAction("复制")
        copy_url_action = menu.addAction("复制链接")
        open_url_action = menu.addAction("在浏览器打开")
        action = menu.exec(self.table.viewport().mapToGlobal(position))
        if action == copy_action:
            item = self.table.itemAt(position)
            if item:
                QApplication.clipboard().setText(item.text())
        elif action == copy_url_action:
            row = self.table.currentRow()
            if row >= 0:
                QApplication.clipboard().setText(self.table.item(row, 2).text())
        elif action == open_url_action:
            row = self.table.currentRow()
            if row >= 0:
                QDesktopServices.openUrl(QUrl(self.table.item(row, 2).text()))

    def _show_about(self):
        QMessageBox.about(self, "关于",
                          "<h2>Supjav 爬虫工具</h2>"
                          "<p>基于 PyQt6 + DrissionPage 的图形化爬虫工具</p>"
                          "<p>功能：自动抓取、去重、本地文件对比、Excel导出</p>"
                          )

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            reply = QMessageBox.question(
                self, "确认", "任务正在运行，确定要退出吗？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.No:
                event.ignore()
                return
            self.worker.stop()
            self.worker.wait(3000)
        self._save_settings()
        event.accept()


def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    # 暗色主题样式表
    app.setStyleSheet(f"""
        QMainWindow {{
            background-color: {DARK_BG};
        }}
        QWidget {{
            background-color: {DARK_BG};
            color: {TEXT_PRIMARY};
        }}
        QGroupBox {{
            font-weight: bold;
            border: 1px solid {BORDER_COLOR};
            border-radius: 8px;
            margin-top: 12px;
            padding-top: 10px;
            background-color: {DARK_CARD};
            color: {TEXT_PRIMARY};
        }}
        QGroupBox::title {{
            subcontrol-origin: margin;
            left: 12px;
            padding: 0 8px;
            color: {ACCENT_BLUE};
        }}
        QLabel#label {{
            color: {TEXT_SECONDARY};
            font-size: 13px;
        }}
        QLineEdit#lineEdit {{
            background-color: {DARK_INPUT};
            border: 1px solid {BORDER_COLOR};
            border-radius: 6px;
            padding: 8px 12px;
            color: {TEXT_PRIMARY};
            font-size: 13px;
            selection-background-color: {ACCENT_BLUE};
        }}
        QLineEdit#lineEdit:focus {{
            border: 1px solid {ACCENT_BLUE};
        }}
        QLineEdit#lineEdit::placeholder {{
            color: {TEXT_SECONDARY};
        }}
        QSpinBox#spinBox {{
            background-color: {DARK_INPUT};
            border: 1px solid {BORDER_COLOR};
            border-radius: 6px;
            padding: 6px;
            color: {TEXT_PRIMARY};
        }}
        QSpinBox#spinBox::up-button, QSpinBox#spinBox::down-button {{
            background-color: {DARK_CARD};
            border: 1px solid {BORDER_COLOR};
        }}
        QCheckBox#checkBox {{
            color: {TEXT_SECONDARY};
            font-size: 13px;
            spacing: 8px;
        }}
        QCheckBox#checkBox::indicator {{
            width: 18px;
            height: 18px;
            border-radius: 4px;
            border: 1px solid {BORDER_COLOR};
            background-color: {DARK_INPUT};
        }}
        QCheckBox#checkBox::indicator:checked {{
            background-color: {ACCENT_GREEN};
            border: 1px solid {ACCENT_GREEN};
        }}
        QPushButton {{
            border: none;
            border-radius: 8px;
            padding: 10px 24px;
            font-size: 14px;
            font-weight: 500;
            color: {DARK_BG};
        }}
        QPushButton#greenBtn {{
            background-color: {ACCENT_GREEN};
        }}
        QPushButton#greenBtn:hover {{
            background-color: #b4f0ae;
        }}
        QPushButton#greenBtn:disabled {{
            background-color: #3d5c3d;
            color: {TEXT_SECONDARY};
        }}
        QPushButton#redBtn {{
            background-color: {ACCENT_RED};
        }}
        QPushButton#redBtn:hover {{
            background-color: #f5a3b8;
        }}
        QPushButton#redBtn:disabled {{
            background-color: #5c3d45;
            color: {TEXT_SECONDARY};
        }}
        QPushButton#blueBtn {{
            background-color: {ACCENT_BLUE};
        }}
        QPushButton#blueBtn:hover {{
            background-color: #a3c8f5;
        }}
        QPushButton#blueBtn:disabled {{
            background-color: #3d4f5c;
            color: {TEXT_SECONDARY};
        }}
        QPushButton#accentBtn {{
            background-color: {ACCENT_PURPLE};
            color: {DARK_BG};
        }}
        QPushButton#accentBtn:hover {{
            background-color: #d4b5f9;
        }}
        QProgressBar#progressBar {{
            border: 1px solid {BORDER_COLOR};
            border-radius: 6px;
            background-color: {DARK_INPUT};
            text-align: center;
            color: {TEXT_PRIMARY};
            font-size: 12px;
            height: 24px;
        }}
        QProgressBar#progressBar::chunk {{
            background-color: {ACCENT_BLUE};
            border-radius: 6px;
        }}
        QTableWidget#dataTable {{
            background-color: {DARK_CARD};
            border: 1px solid {BORDER_COLOR};
            border-radius: 8px;
            gridline-color: {BORDER_COLOR};
            color: {TEXT_PRIMARY};
            font-size: 13px;
            selection-background-color: {ACCENT_BLUE};
            selection-color: {DARK_BG};
        }}
        QTableWidget#dataTable::item {{
            padding: 6px;
            border-bottom: 1px solid {BORDER_COLOR};
        }}
        QTableWidget#dataTable::item:alternate {{
            background-color: #32324a;
        }}
        QTableWidget#dataTable::item:selected {{
            background-color: {ACCENT_BLUE};
            color: {DARK_BG};
        }}
        QHeaderView::section {{
            background-color: {DARK_INPUT};
            color: {TEXT_SECONDARY};
            padding: 10px;
            border: none;
            border-bottom: 2px solid {BORDER_COLOR};
            font-weight: bold;
            font-size: 13px;
        }}
        QTextEdit#logText {{
            background-color: {DARK_CARD};
            border: 1px solid {BORDER_COLOR};
            border-radius: 8px;
            color: {TEXT_PRIMARY};
            font-size: 12px;
            padding: 10px;
            selection-background-color: {ACCENT_BLUE};
        }}
        QMenuBar#menuBar {{
            background-color: {DARK_BG};
            color: {TEXT_PRIMARY};
            border-bottom: 1px solid {BORDER_COLOR};
        }}
        QMenuBar#menuBar::item {{
            background-color: transparent;
            padding: 6px 16px;
        }}
        QMenuBar#menuBar::item:selected {{
            background-color: {DARK_INPUT};
            border-radius: 4px;
        }}
        QMenu#menu {{
            background-color: {DARK_CARD};
            border: 1px solid {BORDER_COLOR};
            border-radius: 6px;
            padding: 6px;
        }}
        QMenu#menu::item {{
            color: {TEXT_PRIMARY};
            padding: 8px 24px;
            border-radius: 4px;
        }}
        QMenu#menu::item:selected {{
            background-color: {ACCENT_BLUE};
            color: {DARK_BG};
        }}
        QMenu#contextMenu {{
            background-color: {DARK_CARD};
            border: 1px solid {BORDER_COLOR};
            border-radius: 6px;
            padding: 6px;
        }}
        QMenu#contextMenu::item {{
            color: {TEXT_PRIMARY};
            padding: 8px 24px;
            border-radius: 4px;
        }}
        QMenu#contextMenu::item:selected {{
            background-color: {ACCENT_BLUE};
            color: {DARK_BG};
        }}
        QStatusBar#statusBar {{
            background-color: {DARK_BG};
            color: {TEXT_SECONDARY};
            border-top: 1px solid {BORDER_COLOR};
        }}
        QSplitter::handle {{
            background-color: {BORDER_COLOR};
        }}
        QMessageBox {{
            background-color: {DARK_CARD};
        }}
        QMessageBox QLabel {{
            color: {TEXT_PRIMARY};
        }}
        QMessageBox QPushButton {{
            background-color: {ACCENT_BLUE};
            color: {DARK_BG};
            padding: 8px 20px;
            border-radius: 6px;
        }}
    """)

    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
