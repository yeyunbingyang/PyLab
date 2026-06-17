'''
============================================================
 PowerPoint 生成: python-pptx
============================================================
对照笔记: 5.python-docx-操作Word和PowerPoint文件.md
  — Presentation() 创建幻灯片
  — slide_layouts[] 选择母版
  — shapes.title / placeholders 编辑内容
  — text_frame.add_paragraph() 添加段落
============================================================
'''

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("python-pptx 未安装")
    print("安装: pip install python-pptx")
    exit(0)

# ═══════════════════════════════════════════════════════════════
# 1. 创建标题页
# ═══════════════════════════════════════════════════════════════
print("=" * 50)
print("【1. 创建标题页】")
print("=" * 50)

prs = Presentation()

# layout[0] = 标题幻灯片
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Welcome to Python"
subtitle.text = "Life is short, I use Python"
print("  标题页已创建")

# ═══════════════════════════════════════════════════════════════
# 2. 内容页 — 带项目符号
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【2. 内容页 — 项目符号】")
print("=" * 50)

# layout[1] = 标题和内容
bullet_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_layout)

shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Introduction'
tf = body_shape.text_frame
tf.text = 'History of Python'

# 一级段落
p = tf.add_paragraph()
p.text = "X'max 1989"
p.level = 1

# 二级段落
p = tf.add_paragraph()
p.text = 'Guido began to write interpreter for Python.'
p.level = 2
print("  内容页已创建 (2个段落层级)")

# ═══════════════════════════════════════════════════════════════
# 3. 自定义内容页
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【3. 自定义内容页】")
print("=" * 50)

slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Python 生态'

tf = slide.placeholders[1].text_frame
for topic in ['Python 基础语法', '标准库 (os/sys/json/csv)', '第三方生态 (requests/openpyxl/PyPDF2)']:
    p = tf.add_paragraph()
    p.text = topic
    p.level = 1
print("  生态页已创建")

# ═══════════════════════════════════════════════════════════════
# 4. 保存
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【4. 保存】")
print("=" * 50)

prs.save('test_pptx.pptx')
print("  已保存 test_pptx.pptx")

print()
print("总结: Presentation → slide_layouts → add_slide → shapes/placeholders → save")
