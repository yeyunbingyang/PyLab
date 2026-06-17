'''
============================================================
 Word 操作: python-docx
============================================================
对照笔记: 5.python-docx-操作Word和PowerPoint文件.md
  — Document() 创建文档
  — add_heading / add_paragraph / add_picture / add_table
  — 样式: bold, font.size, underline
  — 读取段落: doc.paragraphs
  — 模板替换: 占位符 → 真实数据
============================================================
'''

try:
    from docx import Document
    from docx.shared import Cm, Pt, Inches
except ImportError:
    print("python-docx 未安装")
    print("安装: pip install python-docx")
    exit(0)

# ═══════════════════════════════════════════════════════════════
# 1. 创建 Word 文档 — 基础元素
# ═══════════════════════════════════════════════════════════════
print("=" * 50)
print("【1. 创建 Word 文档】")
print("=" * 50)

doc = Document()

# 大标题 (level=0)
doc.add_heading('快快乐乐学Python', 0)

# 段落 + 内联样式
p = doc.add_paragraph('Python是一门非常流行的编程语言，它')
run = p.add_run('简单')
run.bold = True
run.font.size = Pt(18)
p.add_run('而且')
run = p.add_run('优雅')
run.font.size = Pt(18)
run.underline = True
p.add_run('。')

# ═══════════════════════════════════════════════════════════════
# 2. 标题层级 + 列表
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【2. 标题 + 列表】")
print("=" * 50)

doc.add_heading('一级标题', level=1)
doc.add_paragraph('Intense quote', style='Intense Quote')

# 无序列表
doc.add_paragraph('Python 基础语法', style='List Bullet')
doc.add_paragraph('标准库模块', style='List Bullet')
doc.add_paragraph('第三方生态', style='List Bullet')

# 有序列表
doc.add_paragraph('第一步: 安装 Python', style='List Number')
doc.add_paragraph('第二步: 学习基础语法', style='List Number')
doc.add_paragraph('第三步: 实战项目', style='List Number')

print("  已添加 1级标题 + 无序/有序列表")

# ═══════════════════════════════════════════════════════════════
# 3. 表格
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【3. 表格】")
print("=" * 50)

records = (
    ('骆昊', '男', '1995-5-5'),
    ('孙美丽', '女', '1992-2-2')
)

table = doc.add_table(rows=1, cols=3)
table.style = 'Dark List'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '姓名'
hdr_cells[1].text = '性别'
hdr_cells[2].text = '出生日期'

for name, sex, birthday in records:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = sex
    row_cells[2].text = birthday

print("  已添加表格 (3列 x 3行)")

# ═══════════════════════════════════════════════════════════════
# 4. 分页 + 保存
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【4. 分页 + 保存】")
print("=" * 50)

doc.add_page_break()
doc.add_paragraph('这是第二页的内容')
doc.save('demo.docx')
print("  已保存 demo.docx")

# ═══════════════════════════════════════════════════════════════
# 5. 读取 Word 文档
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【5. 读取 Word 文档】")
print("=" * 50)

doc = Document('demo.docx')
print("  段落内容:")
for no, p in enumerate(doc.paragraphs[:6]):
    if p.text.strip():
        print(f"    [{no}] {p.text[:50]}")

# ═══════════════════════════════════════════════════════════════
# 6. 模板替换示例
# ═══════════════════════════════════════════════════════════════
print()
print("=" * 50)
print("【6. 模板替换 — 批量生成】")
print("=" * 50)

print("  用法 (模板文件需提前准备):")
print("    doc = Document('模板.docx')")
print("    for p in doc.paragraphs:")
print("        for run in p.runs:")
print("            if '{name}' in run.text:")
print("                run.text = run.text.replace('{name}', emp['name'])")
print("    doc.save(f'{emp[\"name\"]}离职证明.docx')")

print()
print("总结:")
print("  Document → add_heading / add_paragraph / add_table / add_page_break")
print("  段落读取: doc.paragraphs → p.text / p.runs")
print("  模板替换: 遍历 runs 替换占位符, 保留样式")
